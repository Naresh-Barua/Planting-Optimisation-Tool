import io
from datetime import datetime, timezone

from docx import Document
from docx.shared import RGBColor
from PIL import Image
from pypdf import PdfReader

from src.domains.reporting import FarmReportContract, FarmReportMetadata, RecommendationReportEntry, SaplingReportSummary
from src.services.reporting_export import _get_reason_color_docx, _get_reason_color_pdf, _split_reason, generate_docx, generate_pdf

FARM_METADATA = FarmReportMetadata(
    id=1,
    user_name="Test Farmer",
    rainfall_mm=1200,
    temperature_celsius=25,
    elevation_m=100,
    ph=6.5,
    soil_texture="Loam",
    area_ha=2.5,
    latitude=-9.0,
    longitude=125.0,
)

RECOMMENDATION = RecommendationReportEntry(
    species_id=1,
    species_name="Tectona grandis",
    species_common_name="Teak",
    rank_overall=1,
    score_mcda=0.85,
    key_reasons=["rainfall: inside optimal range"],
)


def _build_report(*, boundary=None, sapling=None, planting_guidance=None) -> FarmReportContract:
    return FarmReportContract(
        farm=FARM_METADATA,
        recommendations=[RECOMMENDATION],
        exclusions=[],
        generated_at=datetime.now(timezone.utc),
        boundary=boundary,
        sapling=sapling,
        planting_guidance=planting_guidance,
    )


def _fake_png_bytes() -> bytes:
    image = Image.new("RGB", (10, 10), color=(34, 139, 34))
    buffer = io.BytesIO()
    image.save(buffer, format="PNG")
    return buffer.getvalue()


def _pdf_text(pdf_bytes: bytes) -> str:
    reader = PdfReader(io.BytesIO(pdf_bytes))
    return "\n".join(page.extract_text() for page in reader.pages)


def _docx_text(docx_bytes: bytes) -> str:
    doc = Document(io.BytesIO(docx_bytes))
    paragraphs = [p.text for p in doc.paragraphs]
    table_cells = [cell.text for table in doc.tables for row in table.rows for cell in row.cells]
    return "\n".join(paragraphs + table_cells)


def _pdf_image_count(pdf_bytes: bytes) -> int:
    reader = PdfReader(io.BytesIO(pdf_bytes))
    return sum(len(page.images) for page in reader.pages)


def _docx_image_count(docx_bytes: bytes) -> int:
    return len(Document(io.BytesIO(docx_bytes)).inline_shapes)


def test_split_reason_with_colon():
    factor, result = _split_reason("rainfall: inside optimal range")
    assert factor == "rainfall"
    assert result == "inside optimal range"


def test_split_reason_without_colon():
    factor, result = _split_reason("suitable rainfall")
    assert factor == ""
    assert result == "suitable rainfall"


def test_split_reason_extra_colons():
    factor, result = _split_reason("rainfall: below minimum: by 200mm")
    assert factor == "rainfall"
    assert result == "below minimum: by 200mm"


def test_split_reason_strips_whitespace():
    factor, result = _split_reason("  ph :  inside range  ")
    assert factor == "ph"
    assert result == "inside range"


def test_get_reason_color_pdf_green():
    assert _get_reason_color_pdf("inside optimal range") == (0, 128, 0)
    assert _get_reason_color_pdf("exact match") == (0, 128, 0)
    assert _get_reason_color_pdf("plateau") == (0, 128, 0)


def test_get_reason_color_pdf_red():
    assert _get_reason_color_pdf("below minimum") == (220, 0, 0)
    assert _get_reason_color_pdf("above maximum") == (220, 0, 0)


def test_get_reason_color_pdf_amber():
    assert _get_reason_color_pdf("marginal") == (200, 120, 0)
    assert _get_reason_color_pdf("unknown result") == (200, 120, 0)


def test_get_reason_color_docx_green():
    assert _get_reason_color_docx("inside optimal range") == RGBColor(0, 128, 0)
    assert _get_reason_color_docx("exact match") == RGBColor(0, 128, 0)
    assert _get_reason_color_docx("plateau") == RGBColor(0, 128, 0)


def test_get_reason_color_docx_red():
    assert _get_reason_color_docx("below minimum") == RGBColor(220, 0, 0)
    assert _get_reason_color_docx("above maximum") == RGBColor(220, 0, 0)


def test_get_reason_color_docx_amber():
    assert _get_reason_color_docx("marginal") == RGBColor(200, 120, 0)
    assert _get_reason_color_docx("unknown result") == RGBColor(200, 120, 0)


def test_generate_pdf_full_report_includes_new_sections():
    sapling = SaplingReportSummary(aligned_count=50, baseline_tree_count=5, additional_sapling_count=45)
    report = _build_report(boundary={"type": "Feature"}, sapling=sapling, planting_guidance="Plant Teak first.")

    pdf_bytes = generate_pdf(report, map_image_bytes=_fake_png_bytes())
    text = _pdf_text(pdf_bytes)

    assert "Farm Boundary Map" in text
    assert "Sapling Capacity" in text
    assert "Existing Trees" in text
    assert "5" in text
    assert "45" in text
    assert "Planting Guidance" in text
    assert "Plant Teak first." in text
    assert "Species Recommendations" in text
    # Not just that the heading rendered: the actual image must be embedded too.
    assert _pdf_image_count(pdf_bytes) == 2  # logo + boundary map


def test_generate_pdf_minimal_report_omits_new_sections():
    report = _build_report(boundary=None, sapling=None, planting_guidance=None)

    pdf_bytes = generate_pdf(report, map_image_bytes=None)
    text = _pdf_text(pdf_bytes)

    assert "Farm Boundary Map" not in text
    assert "Sapling Capacity" not in text
    assert "Planting Guidance" not in text
    assert "Species Recommendations" in text
    assert _pdf_image_count(pdf_bytes) == 1  # logo only, no map


def test_generate_docx_full_report_includes_new_sections():
    sapling = SaplingReportSummary(aligned_count=50, baseline_tree_count=5, additional_sapling_count=45)
    report = _build_report(boundary={"type": "Feature"}, sapling=sapling, planting_guidance="Plant Teak first.")

    docx_bytes = generate_docx(report, map_image_bytes=_fake_png_bytes())
    text = _docx_text(docx_bytes)

    assert "Farm Boundary Map" in text
    assert "Sapling Capacity" in text
    assert "Existing Trees" in text
    assert "5" in text
    assert "45" in text
    assert "Planting Guidance" in text
    assert "Plant Teak first." in text
    assert "Species Recommendations" in text
    # Not just that the heading rendered: the actual image must be embedded too.
    assert _docx_image_count(docx_bytes) == 2  # logo + boundary map


def test_generate_docx_minimal_report_omits_new_sections():
    report = _build_report(boundary=None, sapling=None, planting_guidance=None)

    docx_bytes = generate_docx(report, map_image_bytes=None)
    text = _docx_text(docx_bytes)

    assert "Farm Boundary Map" not in text
    assert "Sapling Capacity" not in text
    assert "Planting Guidance" not in text
    assert "Species Recommendations" in text
    assert _docx_image_count(docx_bytes) == 1  # logo only, no map


def test_generate_pdf_no_recommendations_shows_placeholder_text():
    report = FarmReportContract(
        farm=FARM_METADATA,
        recommendations=[],
        exclusions=[],
        generated_at=datetime.now(timezone.utc),
    )

    text = _pdf_text(generate_pdf(report))

    assert "No recommendations available for this farm." in text


def test_generate_docx_no_recommendations_shows_placeholder_text():
    report = FarmReportContract(
        farm=FARM_METADATA,
        recommendations=[],
        exclusions=[],
        generated_at=datetime.now(timezone.utc),
    )

    text = _docx_text(generate_docx(report))

    assert "No recommendations available for this farm." in text


def test_generate_pdf_with_exclusions_lists_species_and_reasons():
    exclusion = RecommendationReportEntry(
        species_id=2,
        species_name="Excluded Species",
        species_common_name="Excluded Common",
        rank_overall=-1,
        score_mcda=0.0,
        key_reasons=["rainfall: below minimum"],
    )
    report = FarmReportContract(
        farm=FARM_METADATA,
        recommendations=[],
        exclusions=[exclusion],
        generated_at=datetime.now(timezone.utc),
    )

    text = _pdf_text(generate_pdf(report))

    assert "Excluded Species" in text
    assert "Excluded Common" in text
    assert "below minimum" in text


def test_generate_docx_with_exclusions_lists_species_and_reasons():
    exclusion = RecommendationReportEntry(
        species_id=2,
        species_name="Excluded Species",
        species_common_name="Excluded Common",
        rank_overall=-1,
        score_mcda=0.0,
        key_reasons=["rainfall: below minimum"],
    )
    report = FarmReportContract(
        farm=FARM_METADATA,
        recommendations=[],
        exclusions=[exclusion],
        generated_at=datetime.now(timezone.utc),
    )

    text = _docx_text(generate_docx(report))

    assert "Excluded Species" in text
    assert "Excluded Common" in text
    assert "below minimum" in text
